"""
extract_rules.py
================
Stage 2 of the Aetna payer policy rule-intelligence pipeline.

Consumes extraction_rules.json plus the mined signal lexicon (rule_keywords.csv)
and emits one structured row per billing/reimbursement intent into billing_rules.xlsx,
conforming to the reference schema in UHC billing_rules.xlsx.

Sheets generated:
  1. Rules
  2. Code_Index
  3. Signal_Lexicon
  4. Summary

Usage:
    python extract_rules.py
    python extract_rules.py --input extraction_rules.json --keywords rule_keywords.csv --out billing_rules.xlsx --pdf-dir aetna_il_payment_policies
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Set, Tuple

from claim_requirements import assign_claim_columns
from policy_lib import (
    PDFPageIndexer,
    TextUnit,
    clean_guideline_statements,
    expand_range,
    extract_codes,
    extract_conditions,
    extract_exceptions,
    extract_modifiers,
    flatten_policy,
    is_boilerplate,
    load_policies,
    norm_ws,
    split_sentences,
)

# ----------------------------------------------------------------------------
# Signal Lexicon & Scoring
# ----------------------------------------------------------------------------

CATEGORY_WEIGHT = {
    "denial_prohibition": 1.00,
    "reimbursement_action": 1.00,
    "frequency_limit": 0.95,
    "modifier_usage": 0.85,
    "bundling_inclusion": 0.90,
    "place_of_service": 0.80,
    "eligibility_condition": 0.75,
    "coding_requirement": 0.65,
    "documentation": 0.50,
    "general_policy": 0.40,
    "clinical_context": 0.10,
}

DEONTIC_ACTION_RE = re.compile(
    r"\b(will (?:only )?reimburse|reimburs\w*|pay\w*|payable|allow\w*|den\w*|denied|"
    r"report\w*|bill\w*|submit\w*|append\w*|cover\w*|exceed\w*|limit\w*|prohibit\w*|"
    r"exclusive|inclusive|incidental|participat\w*|support\w*)\b",
    re.I
)

STRUCTURAL_ANCHOR_RE = re.compile(
    r"\b(will (?:only )?reimburse|claims (?:that are submitted )?will be denied|"
    r"will be denied when|denied as non-?reimbursable|not separately (?:reported|reimbursed)|"
    r"should (?:not )?be (?:reported|billed)|must be (?:reported|billed)|"
    r"no more than|more than once|maximum unit|highest level of specificity)\b",
    re.I
)


@dataclass
class Signal:
    keyword: str
    category: str
    strength: float
    polarity: str
    role: str


def load_lexicon(path: str) -> Dict[str, Signal]:
    lex: Dict[str, Signal] = {}
    base_dir = os.path.dirname(os.path.abspath(__file__))
    if not os.path.isabs(path) and not os.path.exists(path):
        candidate = os.path.join(base_dir, path)
        if os.path.exists(candidate):
            path = candidate
    if not os.path.exists(path):
        return lex
    with open(path, newline="", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            kw = row.get("keyword", "").strip()
            if kw:
                lex[kw] = Signal(
                    keyword=kw,
                    category=row.get("category", "general_policy"),
                    strength=float(row.get("strength", 0.5) or 0.5),
                    polarity=row.get("polarity", "neutral"),
                    role=row.get("signal_role", "rule_indicator"),
                )
    return lex


_TOK_RE = re.compile(r"<code>|<num>|[a-z][a-z'/\-]+")


def tokenize(sentence: str) -> List[str]:
    s = sentence.lower()
    s = re.sub(r"\b\d{4,5}[tu]?\b", " <code> ", s)
    s = re.sub(r"\b\d+\b", " <num> ", s)
    return _TOK_RE.findall(s)


def sentence_signals(sentence: str, lex: Dict[str, Signal]) -> List[Signal]:
    toks = [t.replace("<code>", "{code}").replace("<num>", "{n}") for t in tokenize(sentence)]
    hits: Dict[str, Signal] = {}
    for n in (1, 2, 3):
        for i in range(len(toks) - n + 1):
            gram = " ".join(toks[i:i + n])
            sig = lex.get(gram)
            if sig and sig.role == "rule_indicator":
                hits[gram] = sig
    return list(hits.values())


def score_sentence(sentence: str, lex: Dict[str, Signal]) -> Tuple[float, List[Signal]]:
    sigs = sentence_signals(sentence, lex)
    best: Dict[str, List[float]] = defaultdict(list)
    for s in sigs:
        best[s.category].append(s.strength * CATEGORY_WEIGHT.get(s.category, 0.4))
    total = 0.0
    for cat, vals in best.items():
        vals.sort(reverse=True)
        total += vals[0] + 0.25 * sum(vals[1:3])
    if STRUCTURAL_ANCHOR_RE.search(sentence):
        total += 0.50
    return round(total, 3), sorted(sigs, key=lambda s: -s.strength)


# ----------------------------------------------------------------------------
# Rule Dataclass
# ----------------------------------------------------------------------------

@dataclass
class Rule:
    rule_id: str
    policy_idx: int
    policy_title: str
    policy_number: Optional[str]
    line_of_business: Optional[str]
    billing_class: Optional[str]
    rule_type: str
    text: str
    codes: List[str] = field(default_factory=list)
    ranges: List[str] = field(default_factory=list)
    modifiers: List[str] = field(default_factory=list)
    conditions: List[str] = field(default_factory=list)
    exceptions: List[str] = field(default_factory=list)
    source: str = ""
    source_page: str = ""
    parent_id: str = ""
    confidence: float = 0.0
    signals: List[str] = field(default_factory=list)
    attachment: str = ""
    section_path: List[str] = field(default_factory=list)
    order: float = 0
    anchor: str = ""
    claim_columns: List[str] = field(default_factory=list)


def classify_rule_type(text: str) -> str:
    t = text.lower()
    if re.search(r"\b(denied|will not reimburse|non-?reimbursable|inappropriate|not separately)\b", t):
        if re.search(r"\b(unless|except|in lieu of|with the exception of|appended with modifier 59)\b", t):
            return "exception"
        return "denial"
    if re.search(r"\b(unless|except|override|in lieu of|instead of)\b", t):
        return "exception"
    if re.search(r"\b(once|per year|per day|maximum unit|no more than|within \d+ days|timeframe|frequency|eight-week|six-month)\b", t):
        return "frequency"
    if re.search(r"\b(incidental|inclusive|integral|component|bundled)\b", t):
        return "bundling"
    if re.search(r"\b(definition|defined as|means\b)\b", t):
        return "definition"
    return "base"


def dedupe(items: Iterable[str]) -> List[str]:
    seen, out = set(), []
    for x in items:
        x_clean = norm_ws(x)
        if x_clean and x_clean not in seen:
            seen.add(x_clean)
            out.append(x_clean)
    return out


# ----------------------------------------------------------------------------
# Specific Code Association per Policy Domain
# ----------------------------------------------------------------------------

def extract_section_e_code_buckets(policy: Dict[str, Any]) -> Dict[str, List[str]]:
    """Parses Section E table rows into labeled buckets."""
    buckets: Dict[str, List[str]] = defaultdict(list)
    rg = policy.get("reimbursement_guidelines") or {}

    current_label = "general"
    for tbl in rg.get("tables", []):
        for row in tbl.get("content", []):
            k = norm_ws(row.get("key"))
            v = norm_ws(row.get("value"))
            full = f"{k} {v}"

            if not k:
                continue

            k_lower = k.lower()
            if "appropriate" in k_lower and "inappropriate" not in k_lower:
                current_label = "appropriate"
            elif "inappropriate" in k_lower:
                current_label = "inappropriate"
            elif "non-covered" in k_lower or "non covered" in k_lower:
                current_label = "non_covered"
            elif "incidental" in k_lower:
                current_label = "incidental"
            elif "planning" in k_lower:
                current_label = "planning"
            elif "consultation" in k_lower:
                current_label = "consultation"
            elif "chiropractic" in k_lower:
                current_label = "chiropractic"

            codes, _ = extract_codes(full)
            if codes:
                buckets[current_label].extend(codes)

    for k in buckets:
        buckets[k] = dedupe(buckets[k])
    return buckets


def assign_codes_to_rule(rule_text: str, policy: Dict[str, Any], buckets: Dict[str, List[str]]) -> List[str]:
    """Binds the most specific clinical codes to an extracted rule."""
    t_lower = rule_text.lower()
    policy_cpts = policy.get("CPT_CODES_USED") or []

    # 1. Direct codes mentioned in rule text
    direct_codes, _ = extract_codes(rule_text)
    if direct_codes:
        return direct_codes

    # 2. Section E labeled buckets
    if "inappropriate" in t_lower and buckets.get("inappropriate"):
        return buckets["inappropriate"]
    if "appropriate" in t_lower and buckets.get("appropriate"):
        return buckets["appropriate"]
    if "non-covered" in t_lower and buckets.get("non_covered"):
        return buckets["non_covered"]
    if "incidental" in t_lower and buckets.get("incidental"):
        return buckets["incidental"]
    if "planning" in t_lower and buckets.get("planning"):
        return buckets["planning"]

    # 3. Domain-specific keyword matching against policy CPTs
    p_title = (policy.get("policy_title") or "").lower()

    if "esketamine" in p_title:
        if "prolonged service" in t_lower:
            return [c for c in policy_cpts if c in ("99415", "99416", "99417")]
        if "inappropriate office" in t_lower or "e&m" in t_lower:
            return [c for c in policy_cpts if c in ("99212", "99213", "99214", "99215")]
        if "appropriate" in t_lower or "dosage" in t_lower:
            return [c for c in policy_cpts if c.startswith("G")]

    if "imrt" in p_title:
        if "planning" in t_lower or "eight-week" in t_lower:
            return ["77301"]
        if "incidental" in t_lower or "two weeks prior" in t_lower:
            return [c for c in policy_cpts if c != "77301"]

    if "chiropractic" in p_title:
        if "non-covered" in t_lower or "98943" in t_lower:
            return ["98943"]
        return [c for c in policy_cpts if c != "98943"]

    if "hydration" in p_title:
        if "fluid" in t_lower or "not separately" in t_lower:
            return [c for c in policy_cpts if c.startswith("J")]
        return [c for c in policy_cpts if c.startswith("96")]

    if "general surgery" in p_title:
        if "sclerosing" in t_lower or "venous" in t_lower:
            return [c for c in policy_cpts if c.startswith("364")]
        if "pilonidal" in t_lower:
            return [c for c in policy_cpts if c.startswith(("100", "117"))]
        if "decompression" in t_lower:
            return [c for c in policy_cpts if c.startswith("63")]

    # If policy has a focused code list (<= 40 codes), inherit policy codes
    if 0 < len(policy_cpts) <= 40:
        return policy_cpts

    return []


# ----------------------------------------------------------------------------
# Cross-Sectional Domain Enrichment
# ----------------------------------------------------------------------------

def enrich_aetna_rule(rule: Rule, policy: Dict[str, Any]) -> None:
    """Enriches rules with domain-specific modifiers, conditions, and codes."""
    title_lower = (policy.get("policy_title") or "").lower()
    text_lower = rule.text.lower()

    if "anatomical modifier" in title_lower:
        rule.modifiers = dedupe(rule.modifiers + [
            "RT", "LT", "50", "E1", "E2", "E3", "E4", "FA", "F1", "F2", "F3", "F4", "F5",
            "F6", "F7", "F8", "F9", "TA", "T1", "T2", "T3", "T4", "T5", "T6", "T7", "T8",
            "T9", "LC", "LD", "LM", "RC", "RI"
        ])
        rule.conditions = dedupe(rule.conditions + [
            "Specific anatomical modifier required (E1-E4 for eyelids, FA/F1-F9 for fingers, TA/T1-T9 for toes, LC/LD/LM/RC/RI for coronary arteries, RT/LT for bilateral sides)",
            "Modifier code must match organ, body part or structure reported by CPT code",
        ])

    elif "never event" in title_lower:
        rule.modifiers = dedupe(rule.modifiers + ["PA", "PB", "PC"])
        rule.conditions = dedupe(rule.conditions + [
            "Modifier PA: Surgery wrong body part",
            "Modifier PB: Surgery wrong patient",
            "Modifier PC: Wrong surgery on patient",
            "Claims billed with Never Event modifiers are denied as non-reimbursable",
        ])

    elif "global surgery" in title_lower:
        if "25" in text_lower or "significant and separately identifiable" in text_lower:
            rule.modifiers = dedupe(rule.modifiers + ["25"])
            rule.conditions = dedupe(rule.conditions + [
                "Modifier 25 required for significant and separately identifiable E&M service on same day as procedure with 0, 10, or 90 global days"
            ])
        if "24" in text_lower or "30 days" in text_lower:
            rule.modifiers = dedupe(rule.modifiers + ["24"])
            rule.conditions = dedupe(rule.conditions + [
                "Modifier 24 required for E&M within 30 days of 30-day surgery with unrelated diagnosis"
            ])

    elif "covid" in title_lower:
        if "59" in text_lower or "separate sample" in text_lower:
            rule.modifiers = dedupe(rule.modifiers + ["59"])
            rule.conditions = dedupe(rule.conditions + [
                "Modifier 59 required to indicate testing of separate sample from same patient on same date of service"
            ])

    elif "place of service" in title_lower:
        if "moderate sedation" in text_lower:
            rule.conditions = dedupe(rule.conditions + [
                "Moderate sedation services by second physician must be billed with facility Place of Service (e.g. POS 19, 21, 22, 24)"
            ])
        if "pos 11" in text_lower:
            rule.conditions = dedupe(rule.conditions + [
                "Services with 'NA' in Non-Facility indicator of CMS fee schedule are denied when rendered in POS 11 (Office)"
            ])

    elif "clinical & correct code editing" in title_lower:
        rule.conditions = dedupe(rule.conditions + [
            "Audited by ClaimsXten and Payment Policy Management software",
            "Illinois Health & Family Services (HFS) state guidelines take precedence",
            "Modifiers 25 and 59 subject to clinical coding validation review",
        ])

    elif "diagnosis code" in title_lower:
        rule.conditions = dedupe(rule.conditions + [
            "ICD-10-CM diagnosis codes must be coded to the highest level of specificity and valid on Date of Service"
        ])


# ----------------------------------------------------------------------------
# Parent Linkage
# ----------------------------------------------------------------------------

def link_parents(rules: List[Rule]) -> None:
    """Connects denial, frequency, bundling, and exception rules to their parent base rule."""
    by_policy: Dict[int, List[Rule]] = defaultdict(list)
    for r in rules:
        by_policy[r.policy_idx].append(r)

    for pol_rules in by_policy.values():
        bases = [r for r in pol_rules if r.rule_type == "base"]
        if not bases:
            continue
        primary_base = bases[0]

        for r in pol_rules:
            if r is primary_base or r.parent_id:
                continue
            if r.rule_type in ("denial", "exception", "frequency", "bundling"):
                r.parent_id = primary_base.rule_id

    by_id = {r.rule_id: r for r in rules}
    for r in rules:
        if r.parent_id and r.parent_id in by_id:
            parent = by_id[r.parent_id]
            note = f"[{r.rule_id}] {r.text[:220]}"
            if note not in parent.exceptions:
                parent.exceptions.append(note)


# ----------------------------------------------------------------------------
# Policy Rule Extraction
# ----------------------------------------------------------------------------

IGNORED_TABLE_KEYS = {
    "medicaid", "medicare", "international statistical classification", "definition",
    "definitions", "provider", "action", "reviewed", "review/revision history", "resources",
    "codes/condition of coverage", "reimbursement guidelines"
}


def extract_policy_rules(
    policy: Dict[str, Any],
    idx: int,
    lex: Dict[str, Signal],
    page_indexer: Optional[PDFPageIndexer] = None
) -> List[Rule]:
    """Extracts all structured billing rules from one policy."""
    title = norm_ws(policy.get("policy_title")) or f"Policy_{idx}"
    pnum = norm_ws(policy.get("policy_number")) or f"P{idx:03d}"
    slug = pnum.replace(" ", "")
    lob = norm_ws(policy.get("Line_of_Business")) or "Illinois Medicaid"
    bclass = norm_ws(policy.get("Billingclass")) or "Professional"

    units = flatten_policy(policy, idx)
    if not units:
        return []

    buckets = extract_section_e_code_buckets(policy)
    counter = [0]

    def next_id() -> str:
        counter[0] += 1
        return f"{slug}-R{counter[0]:03d}"

    rules: List[Rule] = []

    for u in units:
        if u.source_type == "guideline":
            sentences = [u.text]
        elif u.source_type == "table":
            # Filter pure definition or section title rows
            text_lower = u.text.lower()
            if any(text_lower.startswith(ign) for ign in IGNORED_TABLE_KEYS):
                continue
            if len(u.text) < 25 and not u.codes:
                continue
            sentences = [u.text]
        else:
            sentences = [s for s in split_sentences(u.text) if not is_boilerplate(s)]

        for sent in sentences:
            if len(sent) < 20:
                continue

            score, sigs = score_sentence(sent, lex)
            is_rule_like = (
                score >= 0.50 or
                bool(STRUCTURAL_ANCHOR_RE.search(sent)) or
                bool(DEONTIC_ACTION_RE.search(sent)) or
                bool(u.codes)
            )
            if not is_rule_like:
                continue

            r_type = classify_rule_type(sent)
            codes = u.codes if (u.source_type == "table" and u.codes) else assign_codes_to_rule(sent, policy, buckets)
            mods = extract_modifiers(sent)
            conds = extract_conditions(sent)
            excs = extract_exceptions(sent)

            confidence = 0.88
            if codes:
                confidence += 0.05
            if mods:
                confidence += 0.03
            if conds:
                confidence += 0.03
            confidence = min(0.99, round(confidence, 2))

            source_page = ""
            if page_indexer:
                source_page = page_indexer.resolve_page(pnum, title, sent)

            trigger_kws = [s.keyword for s in sigs[:6]]

            rule = Rule(
                rule_id=next_id(),
                policy_idx=idx,
                policy_title=title,
                policy_number=pnum,
                line_of_business=lob,
                billing_class=bclass,
                rule_type=r_type,
                text=norm_ws(sent),
                codes=codes,
                modifiers=mods,
                conditions=dedupe(conds),
                exceptions=dedupe(excs),
                source=u.source_label,
                source_page=source_page,
                confidence=confidence,
                signals=trigger_kws,
                section_path=u.section_path,
                order=u.order * 100 + counter[0],
                anchor=sent,
            )
            enrich_aetna_rule(rule, policy)
            rules.append(rule)

    # Deduplicate rules
    seen: Set[str] = set()
    unique: List[Rule] = []
    for r in sorted(rules, key=lambda r: (-r.confidence, r.order)):
        key = re.sub(r"[^a-z0-9]", "", r.text.lower())
        if key in seen:
            continue
        seen.add(key)
        unique.append(r)

    unique.sort(key=lambda r: r.order)
    return unique


# ----------------------------------------------------------------------------
# Excel Output Generation
# ----------------------------------------------------------------------------

COLUMNS = [
    "Rule_ID", "Policy_Title", "Policy_Number", "Line_of_Business", "Billing_Class",
    "Rule", "CPT_HCPCS_Codes", "Modifiers", "Conditions", "Exceptions",
    "Source", "Source_Page", "Parent_Rule_ID",
    "Rule_Type", "Confidence", "Code_Count", "Trigger_Keywords",
    "Attachment_Reference", "Section_Path", "Required_Claim_Columns"
]

EMPTY_MODIFIER = '""'
CELL_LIMIT = 30000


def cell_val(value: str) -> str:
    v = norm_ws(value)
    return v if len(v) <= CELL_LIMIT else v[:CELL_LIMIT - 3] + "..."


def rule_row(r: Rule) -> Dict[str, Any]:
    codes = list(r.codes)
    code_text = ", ".join(codes)
    if r.ranges:
        code_text = ", ".join(r.ranges) + (f" (expanded: {code_text})" if codes else "")
    return {
        "Rule_ID": r.rule_id,
        "Policy_Title": r.policy_title,
        "Policy_Number": r.policy_number or "",
        "Line_of_Business": r.line_of_business or "",
        "Billing_Class": r.billing_class or "",
        "Rule": cell_val(r.text),
        "CPT_HCPCS_Codes": cell_val(code_text),
        "Modifiers": ", ".join(r.modifiers) if r.modifiers else EMPTY_MODIFIER,
        "Conditions": cell_val(" | ".join(r.conditions)),
        "Exceptions": cell_val(" | ".join(r.exceptions)),
        "Source": r.source,
        "Source_Page": r.source_page,
        "Parent_Rule_ID": r.parent_id,
        "Rule_Type": r.rule_type,
        "Confidence": r.confidence,
        "Code_Count": len(codes),
        "Trigger_Keywords": ", ".join(r.signals),
        "Attachment_Reference": r.attachment,
        "Section_Path": " > ".join(r.section_path),
        "Required_Claim_Columns": ", ".join(r.claim_columns),
    }


def write_workbook(
    rules: List[Rule],
    lexicon_path: str,
    policies: List[Dict[str, Any]],
    out_path: str
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor="1F3864")
    body_font = Font(name="Arial", size=10)

    def style_sheet(ws, widths: Dict[str, int], wrap_cols: Set[str]) -> None:
        for c in ws[1]:
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = body_font
                c.alignment = Alignment(
                    vertical="top",
                    wrap_text=ws.cell(1, c.column).value in wrap_cols
                )
        for i in range(1, ws.max_column + 1):
            name = ws.cell(1, i).value
            ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 16)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    # 1. Sheet: Rules
    ws = wb.active
    ws.title = "Rules"
    ws.append(COLUMNS)
    for r in rules:
        row = rule_row(r)
        ws.append([row[c] for c in COLUMNS])

    style_sheet(
        ws,
        {
            "Rule_ID": 18, "Policy_Title": 36, "Policy_Number": 16, "Line_of_Business": 18,
            "Billing_Class": 14, "Rule": 90, "CPT_HCPCS_Codes": 34, "Modifiers": 14,
            "Conditions": 55, "Exceptions": 45, "Source": 34, "Source_Page": 12,
            "Parent_Rule_ID": 18, "Rule_Type": 15, "Confidence": 12, "Code_Count": 12,
            "Trigger_Keywords": 34, "Attachment_Reference": 24, "Section_Path": 34,
            "Required_Claim_Columns": 46
        },
        {"Rule", "Conditions", "Exceptions", "CPT_HCPCS_Codes", "Source", "Section_Path", "Required_Claim_Columns"}
    )

    # 2. Sheet: Code_Index
    ws2 = wb.create_sheet("Code_Index")
    ws2.append(["Code", "Rule_ID", "Policy_Number", "Policy_Title", "Rule_Type", "Rule_Summary"])
    for r in rules:
        for code in r.codes[:400]:
            ws2.append([
                code,
                r.rule_id,
                r.policy_number or "",
                r.policy_title,
                r.rule_type,
                cell_val(r.text)[:300]
            ])
    style_sheet(
        ws2,
        {"Code": 12, "Rule_ID": 18, "Policy_Number": 16, "Policy_Title": 36, "Rule_Type": 15, "Rule_Summary": 90},
        {"Rule_Summary"}
    )

    # 3. Sheet: Signal_Lexicon
    ws3 = wb.create_sheet("Signal_Lexicon")
    base_dir = os.path.dirname(os.path.abspath(__file__))
    if not os.path.isabs(lexicon_path) and not os.path.exists(lexicon_path):
        candidate = os.path.join(base_dir, lexicon_path)
        if os.path.exists(candidate):
            lexicon_path = candidate
    if os.path.exists(lexicon_path):
        with open(lexicon_path, newline="", encoding="utf-8") as fh:
            reader = csv.reader(fh)
            for row in reader:
                ws3.append(row)
    style_sheet(
        ws3,
        {"keyword": 32, "category": 22, "strength": 12, "strength_band": 14, "polarity": 13, "signal_role": 15},
        set()
    )

    # 4. Sheet: Summary
    ws4 = wb.create_sheet("Summary")
    types = Counter(r.rule_type for r in rules)
    ws4.append(["Metric", "Value"])
    ws4.append(["Policies in source file", len(policies)])
    ws4.append(["Policies yielding rules", len({r.policy_idx for r in rules})])
    ws4.append(["Total rules/intents", len(rules)])
    for t, n in types.most_common():
        ws4.append([f"  of which {t}", n])
    ws4.append(["Rules with CPT/HCPCS codes", sum(1 for r in rules if r.codes)])
    ws4.append(["Rules with modifiers", sum(1 for r in rules if r.modifiers)])
    ws4.append(["Rules with conditions", sum(1 for r in rules if r.conditions)])
    ws4.append(["Rules linked to a parent rule", sum(1 for r in rules if r.parent_id)])
    ws4.append(["Rules enriched from linked PDFs", sum(1 for r in rules if r.attachment)])
    ws4.append(["Distinct codes covered", len({c for r in rules for c in r.codes})])
    ws4.append(["Mean confidence", round(sum(r.confidence for r in rules) / max(len(rules), 1), 3)])
    style_sheet(ws4, {"Metric": 42, "Value": 18}, set())

    wb.save(out_path)


# ----------------------------------------------------------------------------
# CLI Main
# ----------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description="Extract structured Aetna billing rules into Excel")
    ap.add_argument("--input", default="extraction_rules.json", help="Path to input extraction_rules.json")
    ap.add_argument("--keywords", default="rule_keywords.csv", help="Path to rule_keywords.csv")
    ap.add_argument("--out", default="billing_rules.xlsx", help="Path to output Excel workbook")
    ap.add_argument("--json-out", default="billing_rules.json", help="Optional path to output rules JSON")
    ap.add_argument("--pdf-dir", default="aetna_il_payment_policies", help="Folder of source policy PDFs for page indexing")
    args = ap.parse_args()

    base_dir = os.path.dirname(os.path.abspath(__file__))
    if not os.path.isabs(args.input) and not os.path.exists(args.input):
        args.input = os.path.join(base_dir, args.input)
    if not os.path.isabs(args.keywords) and not os.path.exists(args.keywords):
        args.keywords = os.path.join(base_dir, args.keywords)
    if not os.path.isabs(args.out):
        args.out = os.path.join(base_dir, args.out)
    if args.json_out and not os.path.isabs(args.json_out):
        args.json_out = os.path.join(base_dir, args.json_out)
    if args.pdf_dir and not os.path.isabs(args.pdf_dir) and not os.path.exists(args.pdf_dir):
        candidate = os.path.join(base_dir, args.pdf_dir)
        if os.path.exists(candidate):
            args.pdf_dir = candidate

    print(f"Loading lexicon from: {args.keywords}")
    lex = load_lexicon(args.keywords)
    print(f"Loaded {len(lex)} signal keywords.")

    print(f"Loading policies from: {args.input}")
    policies = load_policies(args.input)
    print(f"Loaded {len(policies)} policies.")

    print(f"Indexing source PDFs from: {args.pdf_dir}")
    page_indexer = PDFPageIndexer(args.pdf_dir) if args.pdf_dir and os.path.isdir(args.pdf_dir) else None

    rules: List[Rule] = []
    for idx, pol in enumerate(policies):
        rules.extend(extract_policy_rules(pol, idx, lex, page_indexer))

    link_parents(rules)
    assign_claim_columns(rules)

    print(f"Writing Excel workbook to: {args.out}")
    write_workbook(rules, args.keywords, policies, args.out)

    if args.json_out:
        with open(args.json_out, "w", encoding="utf-8") as fh:
            json.dump([rule_row(r) for r in rules], fh, indent=2, ensure_ascii=False)
        print(f"Wrote JSON rules to: {args.json_out}")

    print("\nExtraction Summary:")
    print(f"  Total rules extracted: {len(rules)}")
    print(f"  Policies covered: {len({r.policy_idx for r in rules})}/{len(policies)}")
    print(f"  Rule types: {dict(Counter(r.rule_type for r in rules))}")
    print(f"  Rules with codes: {sum(1 for r in rules if r.codes)}")
    print(f"  Rules with modifiers: {sum(1 for r in rules if r.modifiers)}")
    print(f"  Rules with conditions: {sum(1 for r in rules if r.conditions)}")
    print(f"  Rules linked to parent: {sum(1 for r in rules if r.parent_id)}")
    print(f"  Distinct codes: {len({c for r in rules for c in r.codes})}")
    print(f"Successfully generated: {args.out}")


if __name__ == "__main__":
    main()
